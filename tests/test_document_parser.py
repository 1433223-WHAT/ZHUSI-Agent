import io
import base64
import json
import threading
import unittest
from http.server import ThreadingHTTPServer
from unittest.mock import patch
from urllib.request import Request, urlopen

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from document_parser import DocumentParseError, parse_document
from server import ArchAIHandler


class DocumentParserTests(unittest.TestCase):
    def test_txt_is_decoded_and_normalized(self):
        result = parse_document("任务书.txt", "建筑面积 3500 平方米\r\n三层以内".encode("utf-8"))
        self.assertEqual("parsed", result["status"])
        self.assertEqual("text/plain", result["mime_type"])
        self.assertIn("建筑面积 3500 平方米\n三层以内", result["text"])

    def test_markdown_is_supported_as_text(self):
        result = parse_document("案例.md", "# 场地分析\n南侧为教学楼".encode("utf-8"))
        self.assertEqual("text/markdown", result["mime_type"])
        self.assertIn("南侧为教学楼", result["text"])

    def test_docx_extracts_paragraphs_and_table_cells(self):
        doc = Document()
        doc.add_paragraph("大学生活动中心任务书")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "总建筑面积"
        table.cell(0, 1).text = "3500㎡"
        stream = io.BytesIO()
        doc.save(stream)
        result = parse_document("任务书.docx", stream.getvalue())
        self.assertEqual("parsed", result["status"])
        self.assertIn("大学生活动中心任务书", result["text"])
        self.assertIn("总建筑面积 | 3500㎡", result["text"])

    @patch("document_parser._parse_legacy_doc", return_value="大学生活动中心\n限高18米")
    def test_legacy_doc_is_supported_when_word_extraction_succeeds(self, _mock_parser):
        result = parse_document("旧任务书.doc", bytes.fromhex("D0CF11E0A1B11AE1") + b"legacy")
        self.assertEqual("parsed", result["status"])
        self.assertEqual("application/msword", result["mime_type"])
        self.assertIn("限高18米", result["text"])

    def test_fake_legacy_doc_signature_is_rejected(self):
        with self.assertRaises(DocumentParseError):
            parse_document("伪装.doc", b"not a legacy word document")

    def test_textless_pdf_reports_ocr_required(self):
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        stream = io.BytesIO()
        writer.write(stream)
        result = parse_document("扫描任务书.pdf", stream.getvalue())
        self.assertEqual("ocr_required", result["status"])
        self.assertEqual(1, result["page_count"])
        self.assertEqual("", result["text"])

    def test_text_pdf_extracts_real_page_content(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        font_reference = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_reference})
        })
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 40 250 Td (Building area 3500 sqm) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        output = io.BytesIO()
        writer.write(output)

        result = parse_document("brief.pdf", output.getvalue())
        self.assertEqual("parsed", result["status"])
        self.assertIn("Building area 3500 sqm", result["text"])

    def test_extension_and_file_signature_must_match(self):
        with self.assertRaises(DocumentParseError):
            parse_document("伪装任务书.pdf", b"this is not a pdf")

    def test_unsupported_type_is_rejected(self):
        with self.assertRaises(DocumentParseError):
            parse_document("模型.dwg", b"data")

    def test_large_file_is_rejected(self):
        with self.assertRaises(DocumentParseError):
            parse_document("过大.txt", b"x" * (10 * 1024 * 1024 + 1))


class DocumentEndpointTests(unittest.TestCase):
    def test_parse_document_endpoint_returns_extracted_text(self):
        server = ThreadingHTTPServer(("127.0.0.1", 0), ArchAIHandler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            payload = json.dumps({
                "filename": "任务书.txt",
                "content_base64": base64.b64encode("限高18米".encode()).decode(),
            }).encode()
            request = Request(
                f"http://127.0.0.1:{server.server_port}/api/parse_document",
                data=payload,
                headers={"Content-Type": "application/json"},
                method="POST",
            )
            with urlopen(request, timeout=5) as response:
                result = json.loads(response.read().decode())
            self.assertEqual("parsed", result["status"])
            self.assertEqual("限高18米", result["text"])
        finally:
            server.shutdown()
            server.server_close()


if __name__ == "__main__":
    unittest.main()
