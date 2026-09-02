# -*- coding: utf-8 -*-
"""把 05_比赛材料/筑思Agent_队友入组资料/*.md 批量转成 docx（保留标题/表格/列表/代码块/加粗）。"""
import sys, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SRC_DIR = Path(r"E:\AI\项目设计\筑智AI_ArchAI\05_比赛材料\筑思Agent_队友入组资料")
INK = RGBColor(0x21, 0x1D, 0x17)
CLAY = RGBColor(0xB3, 0x54, 0x2E)
MUTED = RGBColor(0x75, 0x6C, 0x5E)


def set_font(run, size=11, bold=False, color=None, mono=False):
    run.font.name = "Consolas" if mono else "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_runs(p, text, size=11, base_bold=False, color=None):
    """按 **加粗** 拆分文本为多个 run。"""
    for i, seg in enumerate(re.split(r"\*\*", text)):
        if not seg:
            continue
        r = p.add_run(seg)
        set_font(r, size=size, bold=base_bold or (i % 2 == 1), color=color)
    return p


def convert_md_to_docx(md_path: Path, docx_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            for cl in code_lines:
                p = doc.add_paragraph()
                r = p.add_run(cl if cl else " ")
                set_font(r, size=9, color=MUTED, mono=True)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header):
                cell = t.cell(0, j); cell.text = ""
                add_runs(cell.paragraphs[0], h, size=10, base_bold=True)
            for ri, row in enumerate(rows, 1):
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ""
                    cell = t.cell(ri, j); cell.text = ""
                    add_runs(cell.paragraphs[0], val, size=10)
            doc.add_paragraph()
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            add_runs(p, m.group(2), size={1: 18, 2: 14, 3: 12, 4: 11}.get(level, 11),
                     base_bold=True, color=INK if level <= 2 else None)
            p.paragraph_format.space_before = Pt(12 if level <= 2 else 6)
            p.paragraph_format.space_after = Pt(4)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            r = p.add_run(stripped.lstrip("> ").strip())
            set_font(r, size=11, color=MUTED)
            p.paragraph_format.left_indent = Cm(0.6)
            i += 1
            continue

        # 无序列表
        m = re.match(r"^(\s*)[-*]\s+(.+)$", line)
        if m:
            p = doc.add_paragraph()
            add_runs(p, m.group(2), size=11)
            p.paragraph_format.left_indent = Cm(0.6 + 0.4 * (len(m.group(1)) // 2))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\s*)\d+\.\s+(.+)$", line)
        if m:
            p = doc.add_paragraph()
            add_runs(p, m.group(2), size=11)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 分隔线
        if re.match(r"^[-*_]{3,}$", stripped):
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, stripped, size=11)
        p.paragraph_format.space_after = Pt(3)
        i += 1

    doc.save(docx_path)
    print("saved:", docx_path.name)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    for md in sorted(SRC_DIR.glob("*.md")):
        convert_md_to_docx(md, SRC_DIR / (md.stem + ".docx"))


if __name__ == "__main__":
    main()
