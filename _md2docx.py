# -*- coding: utf-8 -*-
"""Markdown → Word (.docx) 转换（python-docx），用于项目日志。

用法: python _md2docx.py <input.md> <output.docx>
支持: #/##/### 标题、普通段落、``` 代码块、| 表格、- 列表。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def _set_font(run, name="Microsoft YaHei", size=None, color=None, bold=None):
    """设置 run 字体（含中文 eastAsia）。"""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = size
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.font.bold = bold


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    # 默认字体
    style = doc.styles["Normal"]
    _set_font(style, size=Pt(10.5))

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        # 代码块
        if line.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            _set_font(run, name="Consolas", size=Pt(9), color=RGBColor(0x40, 0x40, 0x40))
            p.paragraph_format.left_indent = Pt(12)
            continue
        # 表格
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Table Grid"
            for j, h in enumerate(header):
                table.rows[0].cells[j].text = h
            for r_i, row in enumerate(rows, 1):
                for j, cell in enumerate(row):
                    if j < len(header):
                        table.rows[r_i].cells[j].text = cell
            doc.add_paragraph()
            continue
        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            h = doc.add_heading("", level=min(level, 4))
            run = h.add_run(text)
            if level == 1:
                _set_font(run, size=Pt(18), color=RGBColor(0x8B, 0x45, 0x13))
            elif level == 2:
                _set_font(run, size=Pt(15), color=RGBColor(0x8B, 0x45, 0x13))
            else:
                _set_font(run, size=Pt(12))
            i += 1
            continue
        # 列表
        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(m.group(1), style="List Bullet")
            p.paragraph_format.left_indent = Pt(18)
            i += 1
            continue
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(m.group(1), style="List Number")
            p.paragraph_format.left_indent = Pt(18)
            i += 1
            continue
        # 普通段落（行内加粗/代码去格式化）
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        text = re.sub(r"`(.+?)`", r"\1", text)
        p = doc.add_paragraph(text)
        i += 1

    doc.save(str(docx_path))
    print(f"✅ {docx_path}")


if __name__ == "__main__":
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
