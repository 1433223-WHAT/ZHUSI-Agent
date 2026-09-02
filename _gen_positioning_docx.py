# -*- coding: utf-8 -*-
"""生成《筑思Agent_项目定位_20260818.docx》（中文样式规范）。"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = r"E:\AI\项目设计\筑智AI_ArchAI\05_比赛材料\筑思Agent_项目定位_20260818.docx"

INK = RGBColor(0x21, 0x1D, 0x17)
CLAY = RGBColor(0xB3, 0x54, 0x2E)
MUTED = RGBColor(0x75, 0x6C, 0x5E)


def set_font(run, name_cn="宋体", name_en="Calibri", size=11, bold=False, color=None):
    run.font.name = name_en
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)


def h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, name_cn="黑体", size=18, bold=True, color=INK)
    p.space_after = Pt(6)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, name_cn="黑体", size=14, bold=True, color=CLAY)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p


def para(doc, text, size=11, bold=False, color=None, indent=0, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(3)
    return p


def bullet(doc, text, bold_head=None, size=11):
    p = doc.add_paragraph(style="List Bullet")
    if bold_head:
        r1 = p.add_run(bold_head)
        set_font(r1, size=size, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=size)
    p.paragraph_format.space_after = Pt(2)
    return p


def make_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.cell(0, j)
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_font(r, name_cn="黑体", size=10.5, bold=True)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            set_font(r, size=10.5)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    return t


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    # 标题
    p = doc.add_paragraph()
    r = p.add_run("筑思 Agent · 项目定位")
    set_font(r, name_cn="黑体", size=22, bold=True, color=INK)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2 = doc.add_paragraph()
    r2 = p2.add_run("从「让生成式 AI 参与设计但不替代判断」的学术问题，演进为「证据可追溯的建筑设计协作平台」")
    set_font(r2, size=11, color=MUTED)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3 = doc.add_paragraph()
    r3 = p3.add_run("2026-08-18 · V0.2 · Evidence-grounded Demo")
    set_font(r3, size=10.5, color=MUTED)
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # 一、一句话定位
    h2(doc, "一、一句话定位")
    para(doc, "筑思 Agent 是面向建筑学生的证据约束型设计协作工作台——AI 参与设计的每一步，但每个说法都可追溯到依据，设计决定始终由学生做出。", size=12, bold=True)
    para(doc, "对外一句话：「一个不会替你画图、但会陪你推敲，并且每个说法都有出处的建筑设计 AI 搭档。」", size=11, color=MUTED)

    # 二、目标用户与场景
    h2(doc, "二、目标用户与场景")
    bullet(doc, "建筑专业学生（课程设计 / 竞赛 / 毕业设计）", bold_head="用户：")
    bullet(doc, "方案构思（“我要做教学楼，从哪开始”）、案例查证（“给我看看社区图书馆的案例和图纸”）、图纸分析（上传草图 → 可见事实/推测/未知三层）、评图前推敲（“我的入口太平，怎么改进”）", bold_head="核心场景：")
    bullet(doc, "打开网页 → 新建项目 → 开始设计（无注册门槛，比赛/同学可即开即用）", bold_head="使用形态：")

    # 三、核心价值
    h2(doc, "三、核心价值（为什么存在）")
    bullet(doc, "AI 推进设计（拆解、提问、给动作），但每个设计决定由学生拍板——这是与所有“替你生成方案”的 AI 工具最根本的区别。", bold_head="1. 参与不替代：")
    bullet(doc, "回答引用的每条知识都能点开看来源（案例/理论/方法，来源可追溯或知识库条目），“这句话从哪来”经得起追问。", bold_head="2. 证据可追溯：")
    bullet(doc, "知识库没有依据就不编造——宁可标注“本轮知识库未提供对应依据”，也不伪绑定相似证据。", bold_head="3. 诚实边界：")
    bullet(doc, "案例以实体形式存在（文本 + 平面图/剖面图/照片/分析图），“光之教堂图片”“有平面的吗”“你从知识库里找”这类自然问法随问随取，且上下文可承接。", bold_head="4. 知识可调用：")
    bullet(doc, "项目档案自动沉淀（设计记忆 / 项目文件 / 知识依据 / 过程记录），设计过程本身成为可复盘的成果。", bold_head="5. 过程资产化：")

    # 四、差异化
    h2(doc, "四、差异化（与三类工具划清界限）")
    make_table(doc, ["对比对象", "筑思的不同"],
               [["通用 AI 助手（ChatGPT 等）", "建筑领域知识 + 设计流程边界 + 证据约束：不胡说、不替做决定、回答有出处"],
                ["生成式设计工具（AI 出方案类）", "不生成完整方案替代设计过程，而是推敲方法 + 案例查证 + 判断留白"],
                ["普通知识库 / 检索工具", "案例实体化 + 多模态资产 + 上下文对话式调用，知识是“可对话调用的”，不是“可搜索的”"]],
               widths=[5.2, 10.6])

    # 五、产品边界
    h2(doc, "五、产品边界（明确不做什么）")
    bullet(doc, "不替学生做设计决定（不输出“你应该这样做”的成品方案）")
    bullet(doc, "不编造证据（无依据就标注，不挂相似证据充数）")
    bullet(doc, "不把“建议”当“必须”（输入边界：建议 ≠ 要求）")
    bullet(doc, "不做“只给答案”的问答机（推进边界：一问一瓶颈，把判断留给学生）")

    # 六、定位演进
    h2(doc, "六、定位演进：从初始到现在")
    make_table(doc, ["维度", "初始（2026-08 初）", "现在（2026-08-18）"],
               [["核心问题", "如何让生成式 AI 参与设计但不替代判断", "同一问题，已从“原则”落地为“机制 + 产品”"],
                ["形态", "四层边界原则 + 对话实验", "完整工作台：项目/图纸/案例/依据/设计记忆"],
                ["证据", "Prompt 里写“要区分事实与推测”", "工程机制：Retrieved/Mentioned/Displayed 三层 + 无依据诚实标注"],
                ["知识", "文本检索（25 大师案例）", "案例实体 + 多模态资产 + 上下文检索（+课程题型库扩充中）"],
                ["使用", "本地命令行 / 页面 Demo", "可部署网页（对外演示目标）"],
                ["对手", "通用 AI 的“建筑设计功能”", "通用 AI + 生成式设计工具 + 普通知识库"]],
               widths=[2.6, 6.6, 6.6])

    # 七、不变底线与里程碑
    h2(doc, "七、不变的三条底线与当前里程碑")
    para(doc, "证据边界（不编造）· 决策权（学生主导）· 诚实（无依据就明说）", size=11, bold=True)
    para(doc, "当前里程碑：V0.2 · Evidence-grounded Demo——证据链完整闭环（回答 → 依据 → 追溯 → 诚实无依据标注）+ 知识资产链（案例实体/视觉资产/上下文承接）。正在：课程题型案例库扩充（8 题型 → 30-40 案例）→ 网页产品化 → 外部测试。", size=11)

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
